import csv
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent / 'db-template'
OUT = BASE / 'labores-db.xlsx'
CSV_FILES = [
    'Areas.csv',
    'Supervisors.csv',
    'Variedades.csv',
    'Colaboradores.csv',
    'Formularios.csv',
    'FormularioRows.csv',
    'AssignmentsAudit.csv',
]

wb = Workbook()
# remove default sheet
wb.remove(wb.active)

for fname in CSV_FILES:
    path = BASE / fname
    if not path.exists():
        print('Missing', path)
        continue
    name = Path(fname).stem
    ws = wb.create_sheet(name)
    with path.open(newline='', encoding='utf-8') as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        continue
    # write rows
    for r in rows:
        ws.append(r)
    # create table
    tab = Table(displayName=name, ref=f"A1:{chr(64+len(rows[0]))}{len(rows)}")
    style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    tab.tableStyleInfo = style
    ws.add_table(tab)

wb.save(OUT)
print('Created', OUT)
